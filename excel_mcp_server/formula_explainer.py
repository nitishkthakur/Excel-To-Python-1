"""Formula explainer: translates Excel formulas into human-readable descriptions.

Replaces cell references with column header names and row labels, and
describes what each formula function does in plain English.
"""

from __future__ import annotations

import re
from typing import Any

from openpyxl.utils import get_column_letter, column_index_from_string


# ── Well-known Excel functions and their plain-English templates ──────────

_FUNCTION_DESCRIPTIONS: dict[str, str] = {
    "SUM": "adds up {args}",
    "AVERAGE": "calculates the average of {args}",
    "COUNT": "counts the number of numeric values in {args}",
    "COUNTA": "counts the number of non-empty cells in {args}",
    "COUNTIF": "counts cells in {range} where the value {criteria}",
    "COUNTIFS": "counts cells matching multiple criteria across {args}",
    "SUMIF": "sums values in {sum_range} where {criteria_range} {criteria}",
    "SUMIFS": "sums values matching multiple criteria across {args}",
    "IF": "checks if {condition}; if true, returns {true_val}; otherwise returns {false_val}",
    "IFERROR": "evaluates {expression}; if it results in an error, returns {fallback} instead",
    "VLOOKUP": "looks up {lookup_val} in the first column of {table}, and returns the value from column {col_num}",
    "HLOOKUP": "looks up {lookup_val} in the first row of {table}, and returns the value from row {row_num}",
    "INDEX": "returns the value at a specific position in {array}",
    "MATCH": "finds the position of {lookup_val} in {range}",
    "MAX": "returns the largest value from {args}",
    "MIN": "returns the smallest value from {args}",
    "ABS": "returns the absolute value of {args}",
    "ROUND": "rounds {number} to {digits} decimal places",
    "ROUNDUP": "rounds {number} up to {digits} decimal places",
    "ROUNDDOWN": "rounds {number} down to {digits} decimal places",
    "LEFT": "extracts the first {num_chars} characters from {text}",
    "RIGHT": "extracts the last {num_chars} characters from {text}",
    "MID": "extracts {num_chars} characters from {text} starting at position {start}",
    "LEN": "returns the length of {text}",
    "TRIM": "removes extra spaces from {text}",
    "UPPER": "converts {text} to uppercase",
    "LOWER": "converts {text} to lowercase",
    "CONCATENATE": "joins together {args}",
    "CONCAT": "joins together {args}",
    "TEXT": "formats {value} using the format {format}",
    "TODAY": "returns today's date",
    "NOW": "returns the current date and time",
    "YEAR": "extracts the year from {date}",
    "MONTH": "extracts the month from {date}",
    "DAY": "extracts the day from {date}",
    "DATEDIF": "calculates the difference between {start_date} and {end_date} in {unit}",
    "SUMPRODUCT": "multiplies corresponding values in {args} and sums the results",
    "UNIQUE": "returns unique values from {args}",
    "SORT": "sorts {array}",
    "FILTER": "filters {array} based on {criteria}",
    "XLOOKUP": "looks up {lookup_val} in {lookup_range} and returns the matching value from {return_range}",
    "PERCENTILE": "returns the {k}-th percentile from {array}",
    "MEDIAN": "returns the median of {args}",
    "STDEV": "calculates the standard deviation of {args}",
    "VAR": "calculates the variance of {args}",
    "AND": "returns TRUE if all conditions are true: {args}",
    "OR": "returns TRUE if any condition is true: {args}",
    "NOT": "reverses the logical value of {args}",
}

# Regex to match cell references like A1, $B$2, Sheet1!A1, A1:B10
_CELL_REF_PATTERN = re.compile(
    r"(?:(\w+)!)?"          # optional sheet prefix
    r"(\$?[A-Z]{1,3})"     # column (with optional $)
    r"(\$?\d+)"            # row (with optional $)
)

_RANGE_PATTERN = re.compile(
    r"(?:(\w+)!)?"
    r"(\$?[A-Z]{1,3})(\$?\d+)"
    r":"
    r"(\$?[A-Z]{1,3})(\$?\d+)"
)

_FUNCTION_NAME_PATTERN = re.compile(r"([A-Z][A-Z0-9_.]+)\(")


def explain_formula(
    formula: str,
    headers: list[str],
    current_row: int | None = None,
) -> str:
    """Produce a human-readable explanation of an Excel formula.

    Args:
        formula: The raw formula string (e.g. '=SUM(A2:A10)')
        headers: Column headers from the sheet (index 0 = column A, etc.)
        current_row: The row number where this formula lives.
    """
    if not formula or not formula.startswith("="):
        return formula or ""

    body = formula[1:]  # strip leading '='

    # Build the readable version by substituting references
    readable = _substitute_references(body, headers, current_row)

    # Identify top-level function
    func_match = _FUNCTION_NAME_PATTERN.match(body.strip())
    func_name = func_match.group(1).upper() if func_match else None

    description = _describe_function(func_name, readable, body, headers, current_row)

    parts = []
    if current_row is not None:
        parts.append(f"In row {current_row}: ")
    parts.append(description)
    parts.append(f"  [Raw: {formula}]")

    return "".join(parts)


def _col_to_header(col_str: str, headers: list[str]) -> str:
    """Convert a column letter (e.g. 'B') to its header name."""
    clean = col_str.replace("$", "")
    try:
        idx = column_index_from_string(clean) - 1
        if 0 <= idx < len(headers):
            return headers[idx]
    except (ValueError, AttributeError):
        pass
    return clean


def _cell_to_label(col_str: str, row_str: str, headers: list[str]) -> str:
    """Convert a cell reference like B2 to 'Revenue (row 2)'."""
    header = _col_to_header(col_str, headers)
    row_num = row_str.replace("$", "")
    return f'"{header}" (row {row_num})'


def _range_to_label(
    col1: str, row1: str, col2: str, row2: str, headers: list[str]
) -> str:
    """Convert a range like A2:A10 to '"Name" rows 2-10'."""
    h1 = _col_to_header(col1, headers)
    h2 = _col_to_header(col2, headers)
    r1 = row1.replace("$", "")
    r2 = row2.replace("$", "")
    if h1 == h2:
        return f'"{h1}" (rows {r1} to {r2})'
    return f'"{h1}" (row {r1}) to "{h2}" (row {r2})'


def _substitute_references(
    body: str, headers: list[str], current_row: int | None
) -> str:
    """Replace cell/range references in a formula with human-readable labels."""
    # Replace ranges first (they contain cell refs)
    def range_replacer(m: re.Match) -> str:
        sheet, c1, r1, c2, r2 = m.groups()
        label = _range_to_label(c1, r1, c2, r2, headers)
        if sheet:
            return f"[{sheet}] {label}"
        return label

    result = _RANGE_PATTERN.sub(range_replacer, body)

    # Then replace individual cell references
    def cell_replacer(m: re.Match) -> str:
        sheet, col, row = m.groups()
        label = _cell_to_label(col, row, headers)
        if sheet:
            return f"[{sheet}] {label}"
        return label

    result = _CELL_REF_PATTERN.sub(cell_replacer, result)
    return result


def _describe_function(
    func_name: str | None,
    readable: str,
    raw_body: str,
    headers: list[str],
    current_row: int | None,
) -> str:
    """Generate a natural-language description for a formula."""
    if func_name and func_name in _FUNCTION_DESCRIPTIONS:
        template = _FUNCTION_DESCRIPTIONS[func_name]
        # For simple single-argument templates, fill in readable args
        args_text = _extract_readable_args(readable, func_name)
        filled = template.replace("{args}", args_text)
        filled = filled.replace("{range}", args_text)
        filled = filled.replace("{array}", args_text)
        filled = filled.replace("{expression}", args_text)
        filled = filled.replace("{text}", args_text)
        filled = filled.replace("{number}", args_text)
        filled = filled.replace("{value}", args_text)
        filled = filled.replace("{date}", args_text)
        filled = filled.replace("{lookup_val}", args_text)

        # Multi-arg placeholders
        arg_parts = _split_top_level_args(readable, func_name)
        if len(arg_parts) >= 2:
            filled = filled.replace("{condition}", arg_parts[0])
            filled = filled.replace("{true_val}", arg_parts[1] if len(arg_parts) > 1 else "?")
            filled = filled.replace("{false_val}", arg_parts[2] if len(arg_parts) > 2 else "?")
            filled = filled.replace("{fallback}", arg_parts[1])
            filled = filled.replace("{criteria}", arg_parts[1])
            filled = filled.replace("{criteria_range}", arg_parts[0])
            filled = filled.replace("{sum_range}", arg_parts[2] if len(arg_parts) > 2 else arg_parts[0])
            filled = filled.replace("{table}", arg_parts[1])
            filled = filled.replace("{col_num}", arg_parts[2] if len(arg_parts) > 2 else "?")
            filled = filled.replace("{row_num}", arg_parts[2] if len(arg_parts) > 2 else "?")
            filled = filled.replace("{lookup_range}", arg_parts[1] if len(arg_parts) > 1 else "?")
            filled = filled.replace("{return_range}", arg_parts[2] if len(arg_parts) > 2 else "?")
            filled = filled.replace("{digits}", arg_parts[1])
            filled = filled.replace("{num_chars}", arg_parts[1] if len(arg_parts) > 1 else "?")
            filled = filled.replace("{start}", arg_parts[1] if len(arg_parts) > 1 else "?")
            filled = filled.replace("{format}", arg_parts[1])
            filled = filled.replace("{start_date}", arg_parts[0])
            filled = filled.replace("{end_date}", arg_parts[1] if len(arg_parts) > 1 else "?")
            filled = filled.replace("{unit}", arg_parts[2] if len(arg_parts) > 2 else "?")
            filled = filled.replace("{k}", arg_parts[1] if len(arg_parts) > 1 else "?")

        # Clean up any un-replaced placeholders
        filled = re.sub(r"\{[a-z_]+\}", args_text, filled)

        return f"This formula {filled}."

    # Fallback: arithmetic or unknown function
    if func_name:
        return f"This formula uses {func_name}() to compute: {readable}."
    return f"This formula computes: {readable}."


def _extract_readable_args(readable: str, func_name: str) -> str:
    """Extract the arguments portion from a readable formula string."""
    # Find the opening paren after the function name
    upper = readable.upper()
    idx = upper.find(func_name + "(")
    if idx == -1:
        return readable
    start = idx + len(func_name) + 1
    # Find matching closing paren
    depth = 1
    pos = start
    while pos < len(readable) and depth > 0:
        if readable[pos] == "(":
            depth += 1
        elif readable[pos] == ")":
            depth -= 1
        pos += 1
    return readable[start : pos - 1]


def _split_top_level_args(readable: str, func_name: str) -> list[str]:
    """Split top-level arguments of the function (respecting nested parens)."""
    inner = _extract_readable_args(readable, func_name)
    args: list[str] = []
    depth = 0
    current: list[str] = []
    for ch in inner:
        if ch == "(":
            depth += 1
            current.append(ch)
        elif ch == ")":
            depth -= 1
            current.append(ch)
        elif ch == "," and depth == 0:
            args.append("".join(current).strip())
            current = []
        else:
            current.append(ch)
    if current:
        args.append("".join(current).strip())
    return args
