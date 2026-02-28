"""Range-data retriever — read an arbitrary cell range."""

from __future__ import annotations

from typing import Annotated, Any

from pydantic import Field
from openpyxl.utils import get_column_letter, range_boundaries

from ._base import (
    add_formula_explanations,
    get_headers,
    open_workbook,
    resolve_sheet,
    serialize,
    validate_content_type,
    validate_file,
)


def retrieve(
    file_path: Annotated[str, Field(
        description="Absolute path to the Excel file (.xlsx, .xlsm, .xltx).",
    )],
    *,
    range_ref: Annotated[str, Field(
        description=(
            "Cell range address in Excel notation, e.g. 'B2:E20' or "
            "'A1:D10'. Case-insensitive."
        ),
    )],
    sheet_name: Annotated[str | None, Field(
        default=None,
        description=(
            "Name of the sheet to read from. Defaults to the active "
            "sheet when null."
        ),
    )] = None,
    content_type: Annotated[str, Field(
        default="formulas",
        description=(
            "What cell content to return. Allowed values: "
            "'formulas' (default, **recommended**) — returns raw formulas "
            "with explanations. "
            "'values' — returns computed values only. "
            "'both' — returns computed values plus formula details."
        ),
    )] = "formulas",
    header_row: Annotated[int, Field(
        default=0,
        description=(
            "1-based row number containing column headers. "
            "0 (default) means auto-detect."
        ),
    )] = 0,
) -> dict[str, Any]:
    """Use this retriever when you know the exact cell range you need — e.g.
    from a previous search result or a known layout.

    Returns all cells in the range as a table with column headers resolved
    from the sheet's header row.
    """
    validate_file(file_path)
    validate_content_type(content_type)
    min_col, min_row, max_col_r, max_row_r = range_boundaries(
        range_ref.upper()
    )

    data_only = content_type in ("values", "both")
    wb = open_workbook(file_path, data_only=data_only)
    try:
        ws = resolve_sheet(wb, sheet_name)
        headers, _ = get_headers(ws, header_row=header_row)

        rows: list[dict[str, Any]] = []
        for row_idx in range(min_row, max_row_r + 1):
            row_data: dict[str, Any] = {"_row": row_idx}
            for col_idx in range(min_col, max_col_r + 1):
                header = (
                    headers[col_idx - 1]
                    if col_idx <= len(headers)
                    else get_column_letter(col_idx)
                )
                row_data[header] = serialize(
                    ws.cell(row=row_idx, column=col_idx).value
                )
            rows.append(row_data)

        result: dict[str, Any] = {
            "sheet_name": ws.title,
            "range": range_ref.upper(),
            "rows_returned": len(rows),
            "data": rows,
            "content_type": content_type,
        }
    finally:
        wb.close()

    if content_type in ("formulas", "both"):
        results_list = [result]
        add_formula_explanations(file_path, results_list, sheet_name)
        result = results_list[0]

    return result
