"""Search retriever — find values or formulas across one or all sheets."""

from __future__ import annotations

from typing import Annotated, Any

from pydantic import Field
from openpyxl.utils import get_column_letter

from ._base import (
    get_headers,
    open_workbook,
    resolve_sheets,
    serialize,
    validate_content_type,
    validate_file,
)


def retrieve(
    file_path: Annotated[str, Field(
        description="Absolute path to the Excel file (.xlsx, .xlsm, .xltx).",
    )],
    *,
    query: Annotated[str, Field(
        description=(
            "Substring to search for across cell values or formula "
            "strings. The search is a substring match — the query can "
            "appear anywhere in the cell content."
        ),
    )],
    sheet_name: Annotated[str | None, Field(
        default=None,
        description=(
            "Name of the sheet to search. Pass null to search ALL sheets "
            "in the workbook."
        ),
    )] = None,
    content_type: Annotated[str, Field(
        default="formulas",
        description=(
            "Which cell view to search in. Allowed values: "
            "'formulas' (default) — searches in raw formula strings; "
            "use this to find which cells reference a variable name. "
            "'values' — searches in computed cell values. "
            "'both' — searches both views, deduplicating matches."
        ),
    )] = "formulas",
    header_row: Annotated[int, Field(
        default=0,
        description=(
            "1-based row number containing column headers. "
            "0 (default) means auto-detect."
        ),
    )] = 0,
    case_sensitive: Annotated[bool, Field(
        default=False,
        description=(
            "Whether the search should be case-sensitive. "
            "false (default) performs a case-insensitive match."
        ),
    )] = False,
    max_results: Annotated[int, Field(
        default=50,
        description=(
            "Maximum number of matching cells to return. Default is 50. "
            "Results beyond this limit are truncated."
        ),
    )] = 50,
) -> dict[str, Any]:
    """Use this retriever to locate where a specific variable, label, or
    value appears across the workbook — especially useful for tracing how
    a variable is calculated across multiple sheets.

    Returns matching cells with their sheet name, cell address, column
    header, and value.

    **Tip**: search with ``content_type='formulas'`` to find which cells
    reference a particular variable name in their formulas.
    """
    validate_file(file_path)
    validate_content_type(content_type)

    all_matches: list[dict[str, Any]] = []

    views = []
    if content_type in ("values", "both"):
        views.append(True)   # data_only=True
    if content_type in ("formulas", "both"):
        views.append(False)  # data_only=False

    seen: set[tuple[str, str]] = set()  # (sheet_name, cell_ref)

    for data_only in views:
        if len(all_matches) >= max_results:
            break
        wb = open_workbook(file_path, data_only=data_only)
        try:
            sheets = resolve_sheets(wb, sheet_name)
            for sname, ws in sheets:
                if len(all_matches) >= max_results:
                    break
                _search_sheet(
                    ws, sname, query, header_row,
                    case_sensitive, max_results,
                    all_matches, seen,
                )
        finally:
            wb.close()

    return {
        "query": query,
        "content_type": content_type,
        "match_count": len(all_matches),
        "truncated": len(all_matches) >= max_results,
        "matches": all_matches,
    }


def _search_sheet(
    ws,
    sheet_name: str,
    query: str,
    header_row: int,
    case_sensitive: bool,
    max_results: int,
    matches: list[dict[str, Any]],
    seen: set[tuple[str, str]],
) -> None:
    max_col = ws.max_column or 0
    max_row = ws.max_row or 0
    headers, _ = get_headers(ws, max_col, header_row)
    q = query if case_sensitive else query.lower()

    for row_idx in range(1, max_row + 1):
        if len(matches) >= max_results:
            return
        for col_idx in range(1, max_col + 1):
            if len(matches) >= max_results:
                return
            val = ws.cell(row=row_idx, column=col_idx).value
            if val is None:
                continue
            val_str = str(val)
            compare = val_str if case_sensitive else val_str.lower()
            if q not in compare:
                continue
            cell_ref = f"{get_column_letter(col_idx)}{row_idx}"
            key = (sheet_name, cell_ref)
            if key in seen:
                continue
            seen.add(key)
            header = (
                headers[col_idx - 1]
                if col_idx <= len(headers)
                else get_column_letter(col_idx)
            )
            matches.append({
                "sheet_name": sheet_name,
                "cell": cell_ref,
                "row": row_idx,
                "column": header,
                "value": serialize(val),
            })
