"""Cell-info retriever — deep-dive into a single cell."""

from __future__ import annotations

from typing import Annotated, Any

from pydantic import Field
from openpyxl.utils import coordinate_to_tuple, get_column_letter

from ._base import (
    get_headers,
    open_workbook,
    resolve_sheet,
    serialize,
    validate_content_type,
    validate_file,
)


def retrieve(
    file_path: Annotated[str, Field(
        description="Absolute path to the Excel file (.xlsx, .xlsm, .xltx).",
    )],
    *,
    cell_reference: Annotated[str, Field(
        description=(
            "Cell address to inspect, e.g. 'B5' or 'AA12'. "
            "Case-insensitive — 'b5' and 'B5' are equivalent."
        ),
    )],
    sheet_name: Annotated[str | None, Field(
        default=None,
        description=(
            "Name of the sheet containing the cell. Defaults to the "
            "active sheet when null."
        ),
    )] = None,
    content_type: Annotated[str, Field(
        default="formulas",
        description=(
            "What to return for the cell. Allowed values: "
            "'formulas' (default) — returns the raw formula and a "
            "plain-English explanation if the cell contains a formula; "
            "otherwise returns the cell value. "
            "'values' — returns only the computed value. "
            "'both' — returns the computed value AND the formula with "
            "explanation."
        ),
    )] = "formulas",
    header_row: Annotated[int, Field(
        default=0,
        description=(
            "1-based row number containing column headers. "
            "0 (default) means auto-detect."
        ),
    )] = 0,
) -> dict[str, Any]:
    """Use this retriever to drill into a specific cell after identifying it
    via search or sheet data retrieval.

    Returns the cell's value, data type, column header name, and — if the
    cell contains a formula — the raw formula and a plain-English
    explanation using column headers and row labels.
    """
    validate_file(file_path)
    validate_content_type(content_type)
    row, col = coordinate_to_tuple(cell_reference.upper())

    result: dict[str, Any] = {
        "cell": cell_reference.upper(),
        "row": row,
        "column": col,
        "column_letter": get_column_letter(col),
        "content_type": content_type,
    }

    if content_type in ("values", "both"):
        wb = open_workbook(file_path, data_only=True)
        try:
            ws = resolve_sheet(wb, sheet_name)
            result["sheet_name"] = ws.title
            val = ws.cell(row=row, column=col).value
            headers, _ = get_headers(ws, header_row=header_row)
            result["column_header"] = (
                headers[col - 1] if col <= len(headers) else get_column_letter(col)
            )
            result["value"] = serialize(val)
            result["data_type"] = (
                type(val).__name__ if val is not None else "empty"
            )
        finally:
            wb.close()

    if content_type in ("formulas", "both"):
        wb = open_workbook(file_path, data_only=False)
        try:
            ws = resolve_sheet(wb, sheet_name)
            result["sheet_name"] = ws.title
            cell_val = ws.cell(row=row, column=col).value
            headers, _ = get_headers(ws, header_row=header_row)
            result["column_header"] = (
                headers[col - 1] if col <= len(headers) else get_column_letter(col)
            )
            if isinstance(cell_val, str) and cell_val.startswith("="):
                from ..formula_explainer import explain_formula

                result["formula"] = cell_val
                result["formula_explanation"] = explain_formula(
                    cell_val, headers, row,
                )
            elif content_type == "formulas":
                # No formula — still populate value
                result["value"] = serialize(cell_val)
                result["data_type"] = (
                    type(cell_val).__name__
                    if cell_val is not None
                    else "empty"
                )
        finally:
            wb.close()

    return result
