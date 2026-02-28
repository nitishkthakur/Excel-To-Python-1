"""Formula retriever — extract formulas from one or all sheets."""

from __future__ import annotations

from typing import Any

from ._base import (
    add_formula_explanations,
    extract_formulas,
    get_headers,
    open_workbook,
    package_results,
    resolve_sheets,
    serialize,
    validate_content_type,
    validate_file,
)


def retrieve(
    file_path: str,
    *,
    sheet_name: str | None = None,
    content_type: str = "formulas",
    header_row: int = 0,
) -> dict[str, Any]:
    """Extract all formulas from one or all sheets with explanations.

    Parameters
    ----------
    file_path:    Path to the Excel file.
    sheet_name:   Sheet name, or ``None`` for all sheets.
    content_type: ``"formulas"`` returns formula strings + explanations.
                  ``"values"`` also includes computed values per formula cell.
                  ``"both"`` returns everything.
    header_row:   Row containing headers (0 = auto-detect).
    """
    validate_file(file_path)
    validate_content_type(content_type)

    wb = open_workbook(file_path, data_only=False)
    try:
        sheets = resolve_sheets(wb, sheet_name)
        results = []
        for sname, ws in sheets:
            max_col = ws.max_column or 0
            max_row = ws.max_row or 0
            headers, _ = get_headers(ws, max_col, header_row)
            formulas = extract_formulas(ws, headers, max_col, max_row)
            results.append({
                "sheet_name": sname,
                "formula_count": len(formulas),
                "formulas": formulas,
            })
    finally:
        wb.close()

    # Optionally add computed values for each formula cell
    if content_type in ("values", "both"):
        wb_val = open_workbook(file_path, data_only=True)
        try:
            sheets_val = resolve_sheets(wb_val, sheet_name)
            for i, (_, ws_val) in enumerate(sheets_val):
                if i >= len(results):
                    break
                for fentry in results[i].get("formulas", []):
                    from openpyxl.utils import coordinate_to_tuple
                    row, col = coordinate_to_tuple(fentry["cell"])
                    fentry["computed_value"] = serialize(
                        ws_val.cell(row=row, column=col).value
                    )
        finally:
            wb_val.close()

    return package_results(results, content_type)
