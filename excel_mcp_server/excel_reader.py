"""Core Excel file reading and analysis module.

Handles Excel files from 1KB to 100MB with memory-efficient strategies:
- Files <= 5MB: standard mode (full feature support)
- Files > 5MB: read-only mode (streaming, lower memory)
"""

from __future__ import annotations

import os
import re
from datetime import datetime
from typing import Any

import openpyxl
from openpyxl.utils import get_column_letter


# Threshold in bytes above which we use read_only mode
LARGE_FILE_THRESHOLD = 5 * 1024 * 1024  # 5 MB
MAX_FILE_SIZE = 100 * 1024 * 1024  # 100 MB
DEFAULT_PAGE_SIZE = 100


def list_excel_files(directory: str) -> list[dict[str, Any]]:
    """List all Excel files in a directory with metadata.

    Returns a list of dicts with keys: name, path, size_bytes, modified.
    """
    if not os.path.isdir(directory):
        raise FileNotFoundError(f"Directory not found: {directory}")

    results: list[dict[str, Any]] = []
    for entry in os.scandir(directory):
        if entry.is_file() and entry.name.lower().endswith((".xlsx", ".xlsm", ".xltx")):
            stat = entry.stat()
            results.append({
                "name": entry.name,
                "path": os.path.abspath(entry.path),
                "size_bytes": stat.st_size,
                "size_human": _human_size(stat.st_size),
                "modified": datetime.fromtimestamp(stat.st_mtime).isoformat(),
            })
    results.sort(key=lambda f: f["name"])
    return results


def _human_size(nbytes: int) -> str:
    for unit in ("B", "KB", "MB", "GB"):
        if nbytes < 1024:
            return f"{nbytes:.1f} {unit}"
        nbytes /= 1024
    return f"{nbytes:.1f} TB"


def _is_large_file(file_path: str) -> bool:
    return os.path.getsize(file_path) > LARGE_FILE_THRESHOLD


def _validate_file(file_path: str) -> None:
    if not os.path.isfile(file_path):
        raise FileNotFoundError(f"File not found: {file_path}")
    size = os.path.getsize(file_path)
    if size > MAX_FILE_SIZE:
        raise ValueError(
            f"File too large ({_human_size(size)}). Maximum supported size is "
            f"{_human_size(MAX_FILE_SIZE)}."
        )
    if not file_path.lower().endswith((".xlsx", ".xlsm", ".xltx")):
        raise ValueError("Unsupported file format. Supported: .xlsx, .xlsm, .xltx")


def _open_workbook(
    file_path: str, *, data_only: bool = False, read_only: bool | None = None
) -> openpyxl.Workbook:
    """Open workbook with appropriate mode based on file size."""
    _validate_file(file_path)
    if read_only is None:
        read_only = _is_large_file(file_path)
    return openpyxl.load_workbook(
        file_path, read_only=read_only, data_only=data_only
    )


def get_workbook_info(file_path: str) -> dict[str, Any]:
    """Return high-level metadata about the workbook."""
    _validate_file(file_path)
    stat = os.stat(file_path)
    wb = _open_workbook(file_path, data_only=True)
    try:
        sheets_info = []
        for name in wb.sheetnames:
            ws = wb[name]
            sheets_info.append({
                "name": name,
                "min_row": ws.min_row,
                "max_row": ws.max_row,
                "min_column": ws.min_column,
                "max_column": ws.max_column,
                "dimensions": ws.dimensions if hasattr(ws, "dimensions") else None,
            })
        return {
            "file_path": os.path.abspath(file_path),
            "file_name": os.path.basename(file_path),
            "size_bytes": stat.st_size,
            "size_human": _human_size(stat.st_size),
            "modified": datetime.fromtimestamp(stat.st_mtime).isoformat(),
            "sheet_count": len(wb.sheetnames),
            "sheet_names": wb.sheetnames,
            "sheets": sheets_info,
        }
    finally:
        wb.close()


def get_sheet_names(file_path: str) -> list[str]:
    """Return the list of sheet names."""
    _validate_file(file_path)
    wb = _open_workbook(file_path, data_only=True)
    try:
        return list(wb.sheetnames)
    finally:
        wb.close()


def _resolve_sheet(wb: openpyxl.Workbook, sheet_name: str | None):
    """Return the worksheet; default to the active sheet if name is None."""
    if sheet_name is None:
        return wb.active
    if sheet_name not in wb.sheetnames:
        raise ValueError(
            f"Sheet '{sheet_name}' not found. Available: {wb.sheetnames}"
        )
    return wb[sheet_name]


def _get_headers(ws, max_col: int | None = None) -> list[str]:
    """Read the first row as headers."""
    headers = []
    mc = max_col or ws.max_column or 0
    for col_idx in range(1, mc + 1):
        cell = ws.cell(row=1, column=col_idx)
        val = cell.value
        headers.append(str(val) if val is not None else get_column_letter(col_idx))
    return headers


def get_sheet_summary(
    file_path: str,
    sheet_name: str | None = None,
    detail_level: str = "brief",
) -> dict[str, Any]:
    """Get a summary of a specific sheet.

    detail_level:
        'brief'    - row/column counts, headers, dimensions
        'detailed' - adds data types per column, sample rows, statistics,
                     formula inventory with human-readable explanations
    """
    _validate_file(file_path)
    if detail_level not in ("brief", "detailed"):
        raise ValueError("detail_level must be 'brief' or 'detailed'")

    # For detailed, we need formulas too, so open twice
    wb_data = _open_workbook(file_path, data_only=True)
    try:
        ws_data = _resolve_sheet(wb_data, sheet_name)
        resolved_name = ws_data.title
        max_row = ws_data.max_row or 0
        max_col = ws_data.max_column or 0
        min_row = ws_data.min_row or 1
        headers = _get_headers(ws_data, max_col)

        summary: dict[str, Any] = {
            "sheet_name": resolved_name,
            "total_rows": max(0, max_row - 1),  # exclude header
            "total_columns": max_col,
            "headers": headers,
            "dimensions": ws_data.dimensions if hasattr(ws_data, "dimensions") else None,
            "detail_level": detail_level,
        }

        if detail_level == "detailed":
            # Column data types
            col_types = _analyze_column_types(ws_data, max_col, max_row)
            summary["column_types"] = col_types

            # Sample rows (first 5 data rows)
            sample = _read_rows(ws_data, headers, start_row=2, max_rows=5)
            summary["sample_rows"] = sample

            # Basic statistics for numeric columns
            stats = _compute_statistics(ws_data, headers, max_col, max_row)
            summary["statistics"] = stats

            # Empty/null analysis
            summary["empty_cells"] = _count_empty_cells(ws_data, max_col, max_row)
    finally:
        wb_data.close()

    # Formulas (only for detailed)
    if detail_level == "detailed":
        wb_formula = _open_workbook(file_path, data_only=False)
        try:
            ws_formula = _resolve_sheet(wb_formula, sheet_name)
            formulas = _extract_formulas(ws_formula, headers, max_col, max_row)
            summary["formulas"] = formulas
            summary["formula_count"] = len(formulas)
        finally:
            wb_formula.close()

    return summary


def read_sheet_data(
    file_path: str,
    sheet_name: str | None = None,
    start_row: int = 1,
    max_rows: int = DEFAULT_PAGE_SIZE,
    columns: list[str] | None = None,
) -> dict[str, Any]:
    """Read paginated data from a sheet.

    start_row: 1-based row number (1 = header row).
    max_rows:  number of data rows to return.
    columns:   optional list of column names to include.
    """
    _validate_file(file_path)
    wb = _open_workbook(file_path, data_only=True)
    try:
        ws = _resolve_sheet(wb, sheet_name)
        max_col = ws.max_column or 0
        total_rows = ws.max_row or 0
        headers = _get_headers(ws, max_col)

        # Filter columns
        if columns:
            col_indices = []
            for c in columns:
                if c in headers:
                    col_indices.append(headers.index(c) + 1)  # 1-based
            if not col_indices:
                raise ValueError(f"None of the requested columns found. Available: {headers}")
            selected_headers = [headers[i - 1] for i in col_indices]
        else:
            col_indices = list(range(1, max_col + 1))
            selected_headers = headers

        # Compute actual data start (skip header)
        data_start = max(2, start_row + 1) if start_row == 1 else max(2, start_row)
        data_end = min(total_rows, data_start + max_rows - 1)

        rows: list[dict[str, Any]] = []
        for row_idx in range(data_start, data_end + 1):
            row_data: dict[str, Any] = {"_row": row_idx}
            for ci, col_idx in enumerate(col_indices):
                val = ws.cell(row=row_idx, column=col_idx).value
                row_data[selected_headers[ci]] = _serialize(val)
            rows.append(row_data)

        return {
            "sheet_name": ws.title,
            "headers": selected_headers,
            "start_row": data_start,
            "rows_returned": len(rows),
            "total_data_rows": max(0, total_rows - 1),
            "has_more": data_end < total_rows,
            "next_start_row": data_end + 1 if data_end < total_rows else None,
            "data": rows,
        }
    finally:
        wb.close()


def get_cell_info(
    file_path: str,
    cell_reference: str,
    sheet_name: str | None = None,
) -> dict[str, Any]:
    """Get detailed information about a single cell."""
    _validate_file(file_path)
    from openpyxl.utils import coordinate_to_tuple

    row, col = coordinate_to_tuple(cell_reference.upper())

    # Get computed value
    wb_data = _open_workbook(file_path, data_only=True)
    try:
        ws_data = _resolve_sheet(wb_data, sheet_name)
        cell_data = ws_data.cell(row=row, column=col)
        value = cell_data.value
        headers = _get_headers(ws_data)
        col_header = headers[col - 1] if col <= len(headers) else get_column_letter(col)
    finally:
        wb_data.close()

    # Get formula
    wb_formula = _open_workbook(file_path, data_only=False)
    try:
        ws_formula = _resolve_sheet(wb_formula, sheet_name)
        cell_formula = ws_formula.cell(row=row, column=col)
        raw_formula = None
        if isinstance(cell_formula.value, str) and cell_formula.value.startswith("="):
            raw_formula = cell_formula.value
    finally:
        wb_formula.close()

    result: dict[str, Any] = {
        "cell": cell_reference.upper(),
        "row": row,
        "column": col,
        "column_letter": get_column_letter(col),
        "column_header": col_header,
        "value": _serialize(value),
        "data_type": type(value).__name__ if value is not None else "empty",
    }

    if raw_formula:
        from .formula_explainer import explain_formula
        result["formula"] = raw_formula
        result["formula_explanation"] = explain_formula(raw_formula, headers, row)

    return result


def get_range_data(
    file_path: str,
    range_ref: str,
    sheet_name: str | None = None,
) -> dict[str, Any]:
    """Read data from a specific cell range (e.g. 'A1:D10')."""
    _validate_file(file_path)
    from openpyxl.utils import range_boundaries

    min_col, min_row, max_col_r, max_row_r = range_boundaries(range_ref.upper())

    wb = _open_workbook(file_path, data_only=True)
    try:
        ws = _resolve_sheet(wb, sheet_name)
        headers = _get_headers(ws)

        rows: list[dict[str, Any]] = []
        for row_idx in range(min_row, max_row_r + 1):
            row_data: dict[str, Any] = {"_row": row_idx}
            for col_idx in range(min_col, max_col_r + 1):
                col_letter = get_column_letter(col_idx)
                header = headers[col_idx - 1] if col_idx <= len(headers) else col_letter
                val = ws.cell(row=row_idx, column=col_idx).value
                row_data[header] = _serialize(val)
            rows.append(row_data)

        return {
            "sheet_name": ws.title,
            "range": range_ref.upper(),
            "rows_returned": len(rows),
            "data": rows,
        }
    finally:
        wb.close()


def get_formulas(
    file_path: str,
    sheet_name: str | None = None,
) -> dict[str, Any]:
    """Extract all formulas from a sheet with human-readable explanations."""
    _validate_file(file_path)

    # Get headers from data view
    wb_data = _open_workbook(file_path, data_only=True)
    try:
        ws_data = _resolve_sheet(wb_data, sheet_name)
        max_col = ws_data.max_column or 0
        max_row = ws_data.max_row or 0
        headers = _get_headers(ws_data, max_col)
        resolved_name = ws_data.title
    finally:
        wb_data.close()

    # Get formulas
    wb_formula = _open_workbook(file_path, data_only=False)
    try:
        ws_formula = _resolve_sheet(wb_formula, sheet_name)
        formulas = _extract_formulas(ws_formula, headers, max_col, max_row)
    finally:
        wb_formula.close()

    return {
        "sheet_name": resolved_name,
        "formula_count": len(formulas),
        "formulas": formulas,
    }


def search_in_sheet(
    file_path: str,
    query: str,
    sheet_name: str | None = None,
    case_sensitive: bool = False,
    max_results: int = 50,
) -> dict[str, Any]:
    """Search for a value across all cells in a sheet."""
    _validate_file(file_path)
    wb = _open_workbook(file_path, data_only=True)
    try:
        ws = _resolve_sheet(wb, sheet_name)
        max_col = ws.max_column or 0
        max_row = ws.max_row or 0
        headers = _get_headers(ws, max_col)

        matches: list[dict[str, Any]] = []
        q = query if case_sensitive else query.lower()

        for row_idx in range(1, max_row + 1):
            if len(matches) >= max_results:
                break
            for col_idx in range(1, max_col + 1):
                if len(matches) >= max_results:
                    break
                val = ws.cell(row=row_idx, column=col_idx).value
                if val is None:
                    continue
                val_str = str(val)
                compare = val_str if case_sensitive else val_str.lower()
                if q in compare:
                    header = headers[col_idx - 1] if col_idx <= len(headers) else get_column_letter(col_idx)
                    matches.append({
                        "cell": f"{get_column_letter(col_idx)}{row_idx}",
                        "row": row_idx,
                        "column": header,
                        "value": _serialize(val),
                    })

        return {
            "sheet_name": ws.title,
            "query": query,
            "match_count": len(matches),
            "truncated": len(matches) >= max_results,
            "matches": matches,
        }
    finally:
        wb.close()


def get_sheet_statistics(
    file_path: str,
    sheet_name: str | None = None,
) -> dict[str, Any]:
    """Compute statistics for all numeric columns in a sheet."""
    _validate_file(file_path)
    wb = _open_workbook(file_path, data_only=True)
    try:
        ws = _resolve_sheet(wb, sheet_name)
        max_col = ws.max_column or 0
        max_row = ws.max_row or 0
        headers = _get_headers(ws, max_col)
        stats = _compute_statistics(ws, headers, max_col, max_row)

        return {
            "sheet_name": ws.title,
            "total_data_rows": max(0, max_row - 1),
            "statistics": stats,
        }
    finally:
        wb.close()


def validate_sheet_data(
    file_path: str,
    sheet_name: str | None = None,
) -> dict[str, Any]:
    """Validate data quality: check for empty cells, mixed types, duplicates."""
    _validate_file(file_path)
    wb = _open_workbook(file_path, data_only=True)
    try:
        ws = _resolve_sheet(wb, sheet_name)
        max_col = ws.max_column or 0
        max_row = ws.max_row or 0
        headers = _get_headers(ws, max_col)

        issues: list[dict[str, Any]] = []
        col_types: dict[str, set[str]] = {h: set() for h in headers}
        empty_per_col: dict[str, int] = {h: 0 for h in headers}
        col_values: dict[str, list] = {h: [] for h in headers}

        for row_idx in range(2, max_row + 1):
            for col_idx in range(1, max_col + 1):
                header = headers[col_idx - 1]
                val = ws.cell(row=row_idx, column=col_idx).value
                if val is None or (isinstance(val, str) and val.strip() == ""):
                    empty_per_col[header] += 1
                else:
                    col_types[header].add(type(val).__name__)
                    col_values[header].append(val)

        # Mixed-type columns
        for header, types in col_types.items():
            if len(types) > 1:
                issues.append({
                    "type": "mixed_types",
                    "column": header,
                    "types_found": sorted(types),
                    "severity": "warning",
                })

        # Columns with many empty cells
        total_data = max(0, max_row - 1)
        for header, empty_count in empty_per_col.items():
            if total_data > 0 and empty_count > 0:
                pct = round(100 * empty_count / total_data, 1)
                if pct > 10:
                    issues.append({
                        "type": "missing_values",
                        "column": header,
                        "empty_count": empty_count,
                        "empty_percent": pct,
                        "severity": "warning" if pct < 50 else "error",
                    })

        # Duplicate row detection (sample-based for large files)
        check_rows = min(max_row, 1000)
        seen: set[tuple] = set()
        dup_count = 0
        for row_idx in range(2, check_rows + 1):
            row_tuple = tuple(
                ws.cell(row=row_idx, column=c).value for c in range(1, max_col + 1)
            )
            if row_tuple in seen:
                dup_count += 1
            else:
                seen.add(row_tuple)
        if dup_count > 0:
            issues.append({
                "type": "duplicate_rows",
                "count": dup_count,
                "checked_rows": check_rows - 1,
                "severity": "info",
            })

        return {
            "sheet_name": ws.title,
            "total_data_rows": total_data,
            "total_columns": max_col,
            "issue_count": len(issues),
            "issues": issues,
            "empty_cells_per_column": empty_per_col,
        }
    finally:
        wb.close()


# ── Private helpers ──────────────────────────────────────────────────────

def _serialize(value: Any) -> Any:
    """Convert cell values to JSON-safe types."""
    if value is None:
        return None
    if isinstance(value, datetime):
        return value.isoformat()
    if isinstance(value, (int, float, bool, str)):
        return value
    return str(value)


def _read_rows(
    ws, headers: list[str], start_row: int, max_rows: int
) -> list[dict[str, Any]]:
    max_row = ws.max_row or 0
    max_col = ws.max_column or 0
    end_row = min(max_row, start_row + max_rows - 1)
    rows = []
    for row_idx in range(start_row, end_row + 1):
        row_data: dict[str, Any] = {"_row": row_idx}
        for col_idx in range(1, max_col + 1):
            header = headers[col_idx - 1] if col_idx <= len(headers) else get_column_letter(col_idx)
            row_data[header] = _serialize(ws.cell(row=row_idx, column=col_idx).value)
        rows.append(row_data)
    return rows


def _analyze_column_types(
    ws, max_col: int, max_row: int
) -> dict[str, list[str]]:
    """Determine data types present in each column."""
    headers = _get_headers(ws, max_col)
    col_types: dict[str, set[str]] = {h: set() for h in headers}
    sample_end = min(max_row, 101)  # sample first 100 data rows
    for row_idx in range(2, sample_end + 1):
        for col_idx in range(1, max_col + 1):
            val = ws.cell(row=row_idx, column=col_idx).value
            if val is not None:
                col_types[headers[col_idx - 1]].add(type(val).__name__)
    return {h: sorted(t) for h, t in col_types.items()}


def _compute_statistics(
    ws, headers: list[str], max_col: int, max_row: int
) -> dict[str, dict[str, Any]]:
    """Compute basic statistics for numeric columns."""
    numeric_cols: dict[str, list[float]] = {}
    for col_idx in range(1, max_col + 1):
        vals: list[float] = []
        for row_idx in range(2, max_row + 1):
            v = ws.cell(row=row_idx, column=col_idx).value
            if isinstance(v, (int, float)) and not isinstance(v, bool):
                vals.append(float(v))
        if vals:
            numeric_cols[headers[col_idx - 1]] = vals

    stats = {}
    for header, vals in numeric_cols.items():
        vals_sorted = sorted(vals)
        n = len(vals_sorted)
        total = sum(vals_sorted)
        mean = total / n
        mid = n // 2
        median = (
            vals_sorted[mid]
            if n % 2 == 1
            else (vals_sorted[mid - 1] + vals_sorted[mid]) / 2
        )
        stats[header] = {
            "count": n,
            "sum": round(total, 4),
            "mean": round(mean, 4),
            "median": round(median, 4),
            "min": vals_sorted[0],
            "max": vals_sorted[-1],
        }
    return stats


def _count_empty_cells(
    ws, max_col: int, max_row: int
) -> dict[str, int]:
    headers = _get_headers(ws, max_col)
    empties: dict[str, int] = {h: 0 for h in headers}
    for row_idx in range(2, max_row + 1):
        for col_idx in range(1, max_col + 1):
            val = ws.cell(row=row_idx, column=col_idx).value
            if val is None or (isinstance(val, str) and val.strip() == ""):
                empties[headers[col_idx - 1]] += 1
    return empties


def _extract_formulas(
    ws, headers: list[str], max_col: int, max_row: int
) -> list[dict[str, Any]]:
    """Extract all formulas with explanations."""
    from .formula_explainer import explain_formula

    formulas: list[dict[str, Any]] = []
    for row_idx in range(1, max_row + 1):
        for col_idx in range(1, max_col + 1):
            cell = ws.cell(row=row_idx, column=col_idx)
            val = cell.value
            if isinstance(val, str) and val.startswith("="):
                col_letter = get_column_letter(col_idx)
                col_header = headers[col_idx - 1] if col_idx <= len(headers) else col_letter
                formulas.append({
                    "cell": f"{col_letter}{row_idx}",
                    "row": row_idx,
                    "column_header": col_header,
                    "formula": val,
                    "explanation": explain_formula(val, headers, row_idx),
                })
    return formulas
