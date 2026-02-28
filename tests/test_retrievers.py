"""Comprehensive tests for all Excel retrievers.

Covers every retriever module with tests for:
- Single-sheet and multi-sheet retrieval
- content_type parameter (values, formulas, both)
- header_row parameter (explicit and auto-detect)
- Error handling (missing files, bad sheet names, invalid params)
- Edge cases (empty sheets, unstructured sheets)
"""

from __future__ import annotations

import json
import os

import pytest

from excel_mcp_server.excel_retrievers import (
    cell_info,
    format_output,
    formula,
    get_sheet_names,
    get_workbook_info,
    list_excel_files,
    range_data,
    search,
    sheet_data,
    statistics,
    summary,
    validation,
)
from excel_mcp_server.excel_retrievers._base import (
    detect_header_row,
    get_headers,
    human_size,
    validate_content_type,
    validate_file,
    validate_output_format,
)


# ═════════════════════════════════════════════════════════════════════════
#  _base utilities
# ═════════════════════════════════════════════════════════════════════════


class TestValidateFile:
    def test_missing_file(self):
        with pytest.raises(FileNotFoundError):
            validate_file("/nonexistent/file.xlsx")

    def test_unsupported_format(self, non_excel_file):
        with pytest.raises(ValueError, match="Unsupported"):
            validate_file(non_excel_file)

    def test_valid_file(self, sample_workbook):
        validate_file(sample_workbook)  # should not raise


class TestValidateContentType:
    def test_valid_types(self):
        for ct in ("values", "formulas", "both"):
            validate_content_type(ct)

    def test_invalid_type(self):
        with pytest.raises(ValueError, match="Invalid content_type"):
            validate_content_type("invalid")


class TestValidateOutputFormat:
    def test_valid_formats(self):
        for fmt in ("json", "markdown", "csv"):
            validate_output_format(fmt)

    def test_invalid_format(self):
        with pytest.raises(ValueError, match="Invalid output_format"):
            validate_output_format("xml")


class TestHumanSize:
    def test_bytes(self):
        assert "100.0 B" == human_size(100)

    def test_kilobytes(self):
        assert "1.0 KB" == human_size(1024)

    def test_megabytes(self):
        assert "5.0 MB" == human_size(5 * 1024 * 1024)


class TestDetectHeaderRow:
    def test_standard_sheet(self, sample_workbook):
        import openpyxl

        wb = openpyxl.load_workbook(sample_workbook)
        ws = wb["Sales"]
        assert detect_header_row(ws) == 1
        wb.close()

    def test_unstructured_sheet(self, unstructured_workbook):
        import openpyxl

        wb = openpyxl.load_workbook(unstructured_workbook)
        ws = wb["Messy"]
        row = detect_header_row(ws)
        # Headers "Item", "Amount", "Tax" are in row 3
        assert row == 3
        wb.close()


class TestGetHeaders:
    def test_auto_detect(self, unstructured_workbook):
        import openpyxl

        wb = openpyxl.load_workbook(unstructured_workbook)
        ws = wb["Messy"]
        headers, hrow = get_headers(ws, header_row=0)
        assert hrow == 3
        assert "Item" in headers
        assert "Amount" in headers
        wb.close()

    def test_explicit_row(self, sample_workbook):
        import openpyxl

        wb = openpyxl.load_workbook(sample_workbook)
        ws = wb["Sales"]
        headers, hrow = get_headers(ws, header_row=1)
        assert hrow == 1
        assert headers == ["Product", "Revenue", "Cost", "Profit"]
        wb.close()


# ═════════════════════════════════════════════════════════════════════════
#  Discovery functions
# ═════════════════════════════════════════════════════════════════════════


class TestListExcelFiles:
    def test_lists_xlsx(self, sample_workbook, tmp_dir):
        files = list_excel_files(tmp_dir)
        assert len(files) == 1
        assert files[0]["name"] == "sample.xlsx"
        assert files[0]["size_bytes"] > 0

    def test_missing_directory(self):
        with pytest.raises(FileNotFoundError):
            list_excel_files("/nonexistent/path")

    def test_empty_directory(self, tmp_dir):
        assert list_excel_files(tmp_dir) == []

    def test_multiple_files(self, sample_workbook, empty_workbook, tmp_dir):
        files = list_excel_files(tmp_dir)
        names = {f["name"] for f in files}
        assert "sample.xlsx" in names
        assert "empty.xlsx" in names


class TestGetWorkbookInfo:
    def test_returns_metadata(self, sample_workbook):
        info = get_workbook_info(sample_workbook)
        assert info["sheet_count"] == 3
        assert "Sales" in info["sheet_names"]
        assert "Info" in info["sheet_names"]
        assert "Expenses" in info["sheet_names"]
        assert info["size_bytes"] > 0

    def test_file_not_found(self):
        with pytest.raises(FileNotFoundError):
            get_workbook_info("/no/such/file.xlsx")

    def test_unsupported_format(self, non_excel_file):
        with pytest.raises(ValueError, match="Unsupported"):
            get_workbook_info(non_excel_file)


class TestGetSheetNames:
    def test_returns_all_names(self, sample_workbook):
        names = get_sheet_names(sample_workbook)
        assert names == ["Sales", "Info", "Expenses"]


# ═════════════════════════════════════════════════════════════════════════
#  sheet_data retriever
# ═════════════════════════════════════════════════════════════════════════


class TestSheetData:
    def test_single_sheet_values(self, sample_workbook):
        result = sheet_data.retrieve(
            sample_workbook, sheet_name="Sales", content_type="values",
        )
        assert result["sheet_name"] == "Sales"
        assert result["rows_returned"] == 4
        assert result["headers"] == ["Product", "Revenue", "Cost", "Profit"]
        assert result["content_type"] == "values"

    def test_single_sheet_formulas(self, sample_workbook):
        result = sheet_data.retrieve(
            sample_workbook, sheet_name="Sales", content_type="formulas",
        )
        assert result["content_type"] == "formulas"
        # Should have formula_explanations
        assert "formula_explanations" in result
        assert result["formula_count"] > 0

    def test_single_sheet_both(self, sample_workbook):
        result = sheet_data.retrieve(
            sample_workbook, sheet_name="Sales", content_type="both",
        )
        assert result["content_type"] == "both"
        assert "formula_explanations" in result

    def test_all_sheets(self, sample_workbook):
        result = sheet_data.retrieve(
            sample_workbook, sheet_name=None, content_type="values",
        )
        assert result["sheet_count"] == 3
        assert len(result["sheets"]) == 3
        sheet_names = {s["sheet_name"] for s in result["sheets"]}
        assert sheet_names == {"Sales", "Info", "Expenses"}

    def test_pagination(self, sample_workbook):
        result = sheet_data.retrieve(
            sample_workbook, sheet_name="Sales",
            content_type="values", max_rows=2,
        )
        assert result["rows_returned"] == 2
        assert result["has_more"] is True
        assert result["next_start_row"] is not None

    def test_column_filter(self, sample_workbook):
        result = sheet_data.retrieve(
            sample_workbook, sheet_name="Sales",
            content_type="values", columns=["Product", "Revenue"],
        )
        assert result["headers"] == ["Product", "Revenue"]
        for row in result["data"]:
            assert "Product" in row
            assert "Revenue" in row
            assert "Cost" not in row

    def test_invalid_columns(self, sample_workbook):
        with pytest.raises(ValueError, match="None of the requested"):
            sheet_data.retrieve(
                sample_workbook, sheet_name="Sales",
                content_type="values", columns=["Nonexistent"],
            )

    def test_nonexistent_sheet(self, sample_workbook):
        with pytest.raises(ValueError, match="not found"):
            sheet_data.retrieve(
                sample_workbook, sheet_name="NoSheet",
                content_type="values",
            )

    def test_header_auto_detect(self, unstructured_workbook):
        result = sheet_data.retrieve(
            unstructured_workbook, sheet_name="Messy",
            content_type="values", header_row=0,
        )
        assert result["header_row"] == 3
        assert "Item" in result["headers"]
        assert "Amount" in result["headers"]

    def test_header_explicit(self, unstructured_workbook):
        result = sheet_data.retrieve(
            unstructured_workbook, sheet_name="Messy",
            content_type="values", header_row=3,
        )
        assert result["header_row"] == 3


# ═════════════════════════════════════════════════════════════════════════
#  formula retriever
# ═════════════════════════════════════════════════════════════════════════


class TestFormula:
    def test_single_sheet(self, sample_workbook):
        result = formula.retrieve(
            sample_workbook, sheet_name="Sales",
        )
        assert result["formula_count"] == 6  # D2-D4 + B5, C5, D5
        cells = {f["cell"] for f in result["formulas"]}
        assert "D2" in cells
        assert "B5" in cells

    def test_explanations_use_headers(self, sample_workbook):
        result = formula.retrieve(sample_workbook, sheet_name="Sales")
        sum_f = next(f for f in result["formulas"] if f["cell"] == "B5")
        assert "Revenue" in sum_f["explanation"]

    def test_no_formulas(self, sample_workbook):
        result = formula.retrieve(sample_workbook, sheet_name="Info")
        assert result["formula_count"] == 0

    def test_all_sheets(self, sample_workbook):
        result = formula.retrieve(sample_workbook, sheet_name=None)
        assert result["sheet_count"] == 3
        total = sum(s["formula_count"] for s in result["sheets"])
        assert total >= 7  # Sales: 6 + Expenses: 1

    def test_content_type_values(self, sample_workbook):
        """content_type='values' adds computed_value to each formula."""
        result = formula.retrieve(
            sample_workbook, sheet_name="Sales", content_type="values",
        )
        for f in result["formulas"]:
            assert "computed_value" in f

    def test_content_type_both(self, sample_workbook):
        result = formula.retrieve(
            sample_workbook, sheet_name="Sales", content_type="both",
        )
        for f in result["formulas"]:
            assert "computed_value" in f
            assert "explanation" in f

    def test_unstructured_sheet(self, unstructured_workbook):
        result = formula.retrieve(
            unstructured_workbook, sheet_name="Messy",
        )
        assert result["formula_count"] == 4  # C4, C5, B6, C6


# ═════════════════════════════════════════════════════════════════════════
#  search retriever
# ═════════════════════════════════════════════════════════════════════════


class TestSearch:
    def test_find_text_single_sheet(self, sample_workbook):
        result = search.retrieve(
            sample_workbook, query="Widget", sheet_name="Sales",
            content_type="values",
        )
        assert result["match_count"] == 3

    def test_case_insensitive(self, sample_workbook):
        result = search.retrieve(
            sample_workbook, query="widget", sheet_name="Sales",
            content_type="values", case_sensitive=False,
        )
        assert result["match_count"] == 3

    def test_case_sensitive(self, sample_workbook):
        result = search.retrieve(
            sample_workbook, query="widget", sheet_name="Sales",
            content_type="values", case_sensitive=True,
        )
        assert result["match_count"] == 0

    def test_no_matches(self, sample_workbook):
        result = search.retrieve(
            sample_workbook, query="zzz_nothing", sheet_name="Sales",
            content_type="values",
        )
        assert result["match_count"] == 0

    def test_max_results(self, sample_workbook):
        result = search.retrieve(
            sample_workbook, query="Widget", sheet_name="Sales",
            content_type="values", max_results=2,
        )
        assert result["match_count"] == 2
        assert result["truncated"] is True

    def test_search_all_sheets(self, sample_workbook):
        result = search.retrieve(
            sample_workbook, query="Total", sheet_name=None,
            content_type="values",
        )
        # "Total" appears in Sales row 5 and Expenses row 5
        assert result["match_count"] >= 2
        sheets_found = {m["sheet_name"] for m in result["matches"]}
        assert "Sales" in sheets_found
        assert "Expenses" in sheets_found

    def test_search_formulas(self, sample_workbook):
        result = search.retrieve(
            sample_workbook, query="SUM", sheet_name="Sales",
            content_type="formulas",
        )
        assert result["match_count"] >= 3  # B5, C5, D5

    def test_search_both(self, sample_workbook):
        result = search.retrieve(
            sample_workbook, query="SUM", sheet_name="Sales",
            content_type="both",
        )
        assert result["match_count"] >= 3

    def test_search_includes_sheet_name(self, sample_workbook):
        result = search.retrieve(
            sample_workbook, query="Widget", sheet_name=None,
            content_type="values",
        )
        for match in result["matches"]:
            assert "sheet_name" in match


# ═════════════════════════════════════════════════════════════════════════
#  summary retriever
# ═════════════════════════════════════════════════════════════════════════


class TestSummary:
    def test_brief_single_sheet(self, sample_workbook):
        result = summary.retrieve(
            sample_workbook, sheet_name="Sales",
            content_type="values", detail_level="brief",
        )
        assert result["sheet_name"] == "Sales"
        assert result["total_rows"] == 4
        assert result["total_columns"] == 4
        assert result["headers"] == ["Product", "Revenue", "Cost", "Profit"]
        assert result["detail_level"] == "brief"
        assert "formulas" not in result

    def test_detailed_includes_formulas(self, sample_workbook):
        result = summary.retrieve(
            sample_workbook, sheet_name="Sales",
            content_type="formulas", detail_level="detailed",
        )
        assert result["detail_level"] == "detailed"
        assert "formula_explanations" in result
        assert result["formula_count"] > 0

    def test_detailed_includes_stats(self, sample_workbook):
        result = summary.retrieve(
            sample_workbook, sheet_name="Sales",
            content_type="values", detail_level="detailed",
        )
        assert "statistics" in result
        assert "Revenue" in result["statistics"]
        assert result["statistics"]["Revenue"]["count"] == 3

    def test_detailed_includes_sample_rows(self, sample_workbook):
        result = summary.retrieve(
            sample_workbook, sheet_name="Sales",
            content_type="values", detail_level="detailed",
        )
        assert "sample_rows" in result
        assert len(result["sample_rows"]) > 0
        assert result["sample_rows"][0]["Product"] == "Widget A"

    def test_all_sheets_brief(self, sample_workbook):
        result = summary.retrieve(
            sample_workbook, sheet_name=None,
            content_type="values", detail_level="brief",
        )
        assert result["sheet_count"] == 3
        names = {s["sheet_name"] for s in result["sheets"]}
        assert names == {"Sales", "Info", "Expenses"}

    def test_all_sheets_detailed(self, sample_workbook):
        result = summary.retrieve(
            sample_workbook, sheet_name=None,
            content_type="formulas", detail_level="detailed",
        )
        assert result["sheet_count"] == 3
        for s in result["sheets"]:
            assert "formula_explanations" in s

    def test_invalid_detail_level(self, sample_workbook):
        with pytest.raises(ValueError, match="detail_level"):
            summary.retrieve(
                sample_workbook, sheet_name="Sales",
                detail_level="ultra",
            )

    def test_default_sheet(self, sample_workbook):
        """sheet_name=None retrieves all sheets."""
        result = summary.retrieve(sample_workbook, content_type="values")
        assert result["sheet_count"] == 3
        names = {s["sheet_name"] for s in result["sheets"]}
        assert "Sales" in names

    def test_unstructured_sheet(self, unstructured_workbook):
        result = summary.retrieve(
            unstructured_workbook, sheet_name="Messy",
            content_type="values", detail_level="brief",
        )
        assert result["header_row"] == 3
        assert "Item" in result["headers"]


# ═════════════════════════════════════════════════════════════════════════
#  cell_info retriever
# ═════════════════════════════════════════════════════════════════════════


class TestCellInfo:
    def test_data_cell_values(self, sample_workbook):
        result = cell_info.retrieve(
            sample_workbook, cell_reference="A2",
            sheet_name="Sales", content_type="values",
        )
        assert result["value"] == "Widget A"
        assert result["column_header"] == "Product"
        assert result["data_type"] == "str"

    def test_formula_cell_formulas(self, sample_workbook):
        result = cell_info.retrieve(
            sample_workbook, cell_reference="D2",
            sheet_name="Sales", content_type="formulas",
        )
        assert result["formula"] == "=B2-C2"
        assert "formula_explanation" in result
        assert "Revenue" in result["formula_explanation"]
        assert "Cost" in result["formula_explanation"]

    def test_formula_cell_both(self, sample_workbook):
        result = cell_info.retrieve(
            sample_workbook, cell_reference="D2",
            sheet_name="Sales", content_type="both",
        )
        assert "formula" in result
        assert "value" in result

    def test_numeric_cell(self, sample_workbook):
        result = cell_info.retrieve(
            sample_workbook, cell_reference="B2",
            sheet_name="Sales", content_type="values",
        )
        assert result["value"] == 1000
        assert result["data_type"] == "int"

    def test_non_formula_cell_formulas_mode(self, sample_workbook):
        """A non-formula cell in 'formulas' mode should still show value."""
        result = cell_info.retrieve(
            sample_workbook, cell_reference="A2",
            sheet_name="Sales", content_type="formulas",
        )
        assert result["value"] == "Widget A"
        assert "formula" not in result


# ═════════════════════════════════════════════════════════════════════════
#  range_data retriever
# ═════════════════════════════════════════════════════════════════════════


class TestRangeData:
    def test_read_range_values(self, sample_workbook):
        result = range_data.retrieve(
            sample_workbook, range_ref="A1:B3",
            sheet_name="Sales", content_type="values",
        )
        assert result["rows_returned"] == 3
        # Row 1 is the header row itself
        assert result["data"][0]["Product"] == "Product"
        assert result["data"][1]["Product"] == "Widget A"

    def test_read_range_formulas(self, sample_workbook):
        result = range_data.retrieve(
            sample_workbook, range_ref="A1:D5",
            sheet_name="Sales", content_type="formulas",
        )
        assert "formula_explanations" in result
        assert result["formula_count"] > 0

    def test_read_range_both(self, sample_workbook):
        result = range_data.retrieve(
            sample_workbook, range_ref="A1:D5",
            sheet_name="Sales", content_type="both",
        )
        assert "formula_explanations" in result


# ═════════════════════════════════════════════════════════════════════════
#  statistics retriever
# ═════════════════════════════════════════════════════════════════════════


class TestStatistics:
    def test_single_sheet(self, sample_workbook):
        result = statistics.retrieve(
            sample_workbook, sheet_name="Sales",
        )
        stats = result["statistics"]
        assert "Revenue" in stats
        assert stats["Revenue"]["count"] == 3
        assert stats["Revenue"]["sum"] == 4300
        assert stats["Revenue"]["min"] == 800
        assert stats["Revenue"]["max"] == 2500

    def test_all_sheets(self, sample_workbook):
        result = statistics.retrieve(sample_workbook, sheet_name=None)
        assert result["sheet_count"] == 3

    def test_empty_sheet(self, empty_workbook):
        result = statistics.retrieve(empty_workbook)
        assert result["statistics"] == {}


# ═════════════════════════════════════════════════════════════════════════
#  validation retriever
# ═════════════════════════════════════════════════════════════════════════


class TestValidation:
    def test_single_sheet(self, sample_workbook):
        result = validation.retrieve(
            sample_workbook, sheet_name="Sales",
        )
        assert "issues" in result
        assert result["total_data_rows"] == 4

    def test_mixed_types(self, mixed_type_workbook):
        result = validation.retrieve(mixed_type_workbook)
        type_issues = [i for i in result["issues"] if i["type"] == "mixed_types"]
        assert len(type_issues) == 1
        assert type_issues[0]["column"] == "Value"

    def test_all_sheets(self, sample_workbook):
        result = validation.retrieve(sample_workbook, sheet_name=None)
        assert result["sheet_count"] == 3

    def test_empty_cells_reported(self, sample_workbook):
        result = validation.retrieve(
            sample_workbook, sheet_name="Sales",
        )
        assert isinstance(result["empty_cells_per_column"], dict)


# ═════════════════════════════════════════════════════════════════════════
#  Multi-sheet comprehensive tests
# ═════════════════════════════════════════════════════════════════════════


class TestMultiSheet:
    """End-to-end tests for multi-sheet operations."""

    def test_formulas_across_all_sheets(self, sample_workbook):
        result = formula.retrieve(sample_workbook, sheet_name=None)
        assert result["sheet_count"] == 3
        # Collect all formulas
        all_formulas = []
        for s in result["sheets"]:
            all_formulas.extend(s["formulas"])
        assert len(all_formulas) >= 7

    def test_search_across_all_sheets(self, sample_workbook):
        result = search.retrieve(
            sample_workbook, query="500",
            sheet_name=None, content_type="values",
        )
        # "500" appears in Expenses (Travel=500)
        assert result["match_count"] >= 1

    def test_summary_all_sheets_detailed(self, sample_workbook):
        result = summary.retrieve(
            sample_workbook, sheet_name=None,
            content_type="formulas", detail_level="detailed",
        )
        assert result["sheet_count"] == 3
        for s in result["sheets"]:
            assert "headers" in s
            assert "formula_explanations" in s

    def test_different_retrievers_per_sheet(self, sample_workbook):
        """Simulate the LLM using different retrievers for different sheets."""
        # Summary for Sales
        sales_summary = summary.retrieve(
            sample_workbook, sheet_name="Sales",
            content_type="formulas", detail_level="detailed",
        )
        assert sales_summary["sheet_name"] == "Sales"

        # Formulas for Expenses
        expenses_formulas = formula.retrieve(
            sample_workbook, sheet_name="Expenses",
        )
        assert expenses_formulas["sheet_name"] == "Expenses"

        # Data for Info
        info_data = sheet_data.retrieve(
            sample_workbook, sheet_name="Info", content_type="values",
        )
        assert info_data["sheet_name"] == "Info"
