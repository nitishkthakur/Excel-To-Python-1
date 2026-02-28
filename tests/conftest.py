"""Shared test fixtures for Excel MCP Server tests."""

from __future__ import annotations

import os
import tempfile

import openpyxl
import pytest


@pytest.fixture()
def tmp_dir():
    """Provide a temporary directory that is cleaned up after the test."""
    with tempfile.TemporaryDirectory() as d:
        yield d


@pytest.fixture()
def sample_workbook(tmp_dir):
    """Create a small sample workbook with data and formulas.

    Sheet 'Sales':
        A          B         C        D
    1   Product    Revenue   Cost     Profit
    2   Widget A   1000      400      =B2-C2
    3   Widget B   2500      1100     =B3-C3
    4   Widget C   800       350      =B4-C4
    5   Total      =SUM(B2:B4)  =SUM(C2:C4)  =SUM(D2:D4)

    Sheet 'Info':
        A          B
    1   Key        Value
    2   Year       2024
    3   Region     North
    """
    path = os.path.join(tmp_dir, "sample.xlsx")
    wb = openpyxl.Workbook()

    # --- Sales sheet ---
    ws = wb.active
    ws.title = "Sales"
    ws.append(["Product", "Revenue", "Cost", "Profit"])
    ws.append(["Widget A", 1000, 400, None])
    ws["D2"] = "=B2-C2"
    ws.append(["Widget B", 2500, 1100, None])
    ws["D3"] = "=B3-C3"
    ws.append(["Widget C", 800, 350, None])
    ws["D4"] = "=B4-C4"
    ws.append(["Total", None, None, None])
    ws["B5"] = "=SUM(B2:B4)"
    ws["C5"] = "=SUM(C2:C4)"
    ws["D5"] = "=SUM(D2:D4)"

    # --- Info sheet ---
    ws2 = wb.create_sheet("Info")
    ws2.append(["Key", "Value"])
    ws2.append(["Year", 2024])
    ws2.append(["Region", "North"])

    wb.save(path)
    wb.close()
    return path


@pytest.fixture()
def empty_workbook(tmp_dir):
    """An Excel file with a single empty sheet."""
    path = os.path.join(tmp_dir, "empty.xlsx")
    wb = openpyxl.Workbook()
    wb.save(path)
    wb.close()
    return path


@pytest.fixture()
def non_excel_file(tmp_dir):
    """A plain text file that is not an Excel file."""
    path = os.path.join(tmp_dir, "data.txt")
    with open(path, "w") as f:
        f.write("not excel")
    return path
